# -*- coding: utf-8 -*-
"""Extract plain text (and optionally embedded figures) from source materials so
note-writer agents can ground their notes in real content.

Usage:
    python extract_sources.py pptx  "deck.pptx"          [out.txt]
    python extract_sources.py docx   "report.docx"        [out.txt]
    python extract_sources.py figures "report.docx" out_dir   # dump embedded images
"""
import sys, os


def pptx_text(path, out):
    from pptx import Presentation
    p = Presentation(path)
    count = 0
    with open(out, "w", encoding="utf-8") as f:
        for i, s in enumerate(p.slides, 1):
            count = i
            parts = [sh.text_frame.text.strip().replace("\n", " | ")
                     for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            f.write(f"### Slide {i}\n{'  ||  '.join(parts)}\n\n")
    print(f"wrote {out} ({count} slides)")


def docx_text(path, out):
    from docx import Document
    d = Document(path)
    with open(out, "w", encoding="utf-8") as f:
        for para in d.paragraphs:
            if para.text.strip():
                f.write(para.text.strip() + "\n")
    print(f"wrote {out} ({len(d.paragraphs)} paragraphs)")


def docx_figures(path, out_dir):
    from docx import Document
    from docx.oxml.ns import qn
    d = Document(path); os.makedirs(out_dir, exist_ok=True)
    n = 0
    for el in d.element.body.iter():
        if el.tag.split('}')[-1] == 'blip':
            rid = el.get(qn('r:embed'))
            if not rid:
                continue
            try:
                part = d.part.related_parts[rid]
            except KeyError:
                continue
            n += 1
            ext = os.path.splitext(part.partname)[1] or ".png"
            with open(os.path.join(out_dir, f"fig_{n:03d}{ext}"), "wb") as f:
                f.write(part.blob)
    print(f"extracted {n} figures to {out_dir}")


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    mode, path = sys.argv[1], sys.argv[2]
    if mode == "pptx":
        pptx_text(path, sys.argv[3] if len(sys.argv) > 3 else "pptx_full.txt")
    elif mode == "docx":
        docx_text(path, sys.argv[3] if len(sys.argv) > 3 else "report_full.txt")
    elif mode == "figures":
        docx_figures(path, sys.argv[3] if len(sys.argv) > 3 else "report_figs")
    else:
        sys.exit(__doc__)


if __name__ == "__main__":
    main()
